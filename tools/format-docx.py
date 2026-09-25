# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx==1.2.0"]
# ///

import os
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


for entrada in os.environ.get("QUARTO_PROJECT_OUTPUT_FILES", "").splitlines():
    caminho = Path(entrada)
    if caminho.suffix.lower() != ".docx":
        continue
    documento = Document(caminho)
    corpo = documento.element.body
    for tabela in list(corpo.findall(qn("w:tbl"))):
        celulas = tabela.findall("./" + qn("w:tr") + "/" + qn("w:tc"))
        if len(celulas) != 1:
            continue
        celula = celulas[0]
        legendas = [p for p in celula.findall(qn("w:p"))
                    if any(s.get(qn("w:val")) in ("ImageCaption", "TableCaption")
                           for s in p.iter(qn("w:pStyle")))]
        if len(legendas) != 1:
            continue
        legenda = legendas[0]
        tabela_aninhada = celula.find(qn("w:tbl"))
        # O invólucro de float de uma célula do Quarto herda as bordas da tabela de dados e perde o alinhamento.
        for propriedades in list(legenda.findall(qn("w:pPr"))):
            legenda.remove(propriedades)
        paragrafo = Paragraph(legenda, documento)
        paragrafo.style = "Table Caption" if tabela_aninhada is not None else "Image Caption"
        paragrafo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragrafo.paragraph_format.keep_with_next = True
        for texto in legenda.iter(qn("w:t")):
            if texto.text:
                texto.text = re.sub(r"^((?:Tabela|Figura)\s+[A-Z0-9]+)\s*[-:]\s*", r"\1 - ", texto.text)
        for filho in list(celula):
            if filho.tag == qn("w:tcPr"):
                continue
            if filho.tag == qn("w:p") and len(filho) == 0:
                continue
            tabela.addprevious(filho)
        corpo.remove(tabela)

    for tabela in documento.tables:
        fonte = tabela._tbl.getnext()
        while fonte is not None and fonte.tag in (qn("w:bookmarkStart"), qn("w:bookmarkEnd")):
            fonte = fonte.getnext()
        notas = fonte.getnext() if fonte is not None else None
        if (fonte is not None and notas is not None
                and fonte.tag == qn("w:p") and notas.tag == qn("w:p")
                and Paragraph(fonte, documento).style.name == "Fonte"
                and Paragraph(notas, documento).style.name == "Notas"):
            linhas_de_dados = len(tabela.rows)
            rodape = tabela.add_row().cells
            celula = rodape[0].merge(rodape[-1])
            for filho in list(celula._tc):
                if filho.tag != qn("w:tcPr"):
                    celula._tc.remove(filho)
            celula._tc.append(fonte)
            celula._tc.append(notas)
            bordas = OxmlElement("w:tcBorders")
            for lado in ("top", "bottom", "left", "right"):
                borda = OxmlElement(f"w:{lado}")
                borda.set(qn("w:val"), "single" if lado == "top" else "nil")
                borda.set(qn("w:sz"), "8")
                borda.set(qn("w:color"), "000000")
                bordas.append(borda)
            celula._tc.get_or_add_tcPr().append(bordas)
            bordas_tabela = OxmlElement("w:tblBorders")
            inferior = OxmlElement("w:bottom")
            inferior.set(qn("w:val"), "nil")
            bordas_tabela.append(inferior)
            tabela._tbl.tblPr.append(bordas_tabela)
            # O Writer não mantém a última linha de uma tabela junto de um parágrafo de fonte externo.
            for linha in tabela.rows[:-1] if linhas_de_dados < 12 else tabela.rows[-2:-1]:
                for celula_da_linha in linha.cells:
                    for paragrafo in celula_da_linha.paragraphs:
                        paragrafo.paragraph_format.keep_with_next = True
            Paragraph(notas, documento).paragraph_format.keep_with_next = False
            Paragraph(notas, documento).paragraph_format.keep_together = True
        for linha in tabela.rows:
            propriedades = linha._tr.get_or_add_trPr()
            if propriedades.find(qn("w:cantSplit")) is None:
                propriedades.append(OxmlElement("w:cantSplit"))
            for celula in linha.cells:
                for paragrafo in celula.paragraphs:
                    propriedades = paragrafo._p.find(qn("w:pPr"))
                    if propriedades is not None:
                        # O Quarto acrescenta o alinhamento central depois do alinhamento explícito da pipe table.
                        for alinhamento in propriedades.findall(qn("w:jc"))[1:]:
                            propriedades.remove(alinhamento)
    for paragrafo in documento.paragraphs:
        if paragrafo.style.name == "Notas":
            paragrafo.paragraph_format.keep_together = True
        if paragrafo._p.find(".//" + qn("w:drawing")) is not None:
            paragrafo.paragraph_format.keep_with_next = True
    documento.save(caminho)
    print(f"formatado {caminho}")
