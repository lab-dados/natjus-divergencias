# /// script
# requires-python = ">=3.10"
# dependencies = ["docxcompose==1.4.0", "python-docx==1.2.0", "setuptools<81"]
# ///
"""Junta o artigo e o apêndice num arquivo Word só.

Uso: uv run tools/compose-docx.py paper.docx appendix.docx saida.docx

O Quarto numera tabelas e figuras com um esquema só por documento, e o
apêndice precisa de letras onde o artigo usa algarismos; por isso os dois são
renderizados à parte e juntados aqui. O docxcompose leva imagens, hyperlinks e
notas de rodapé com ids de relacionamento novos; o apêndice começa em página
nova. O setuptools fica fixado porque o docxcompose importa pkg_resources, que
o setuptools retirou na versão 81.
"""
import sys

from docx import Document
from docx.enum.text import WD_BREAK
from docxcompose.composer import Composer

caminho_artigo, caminho_apendice, caminho_saida = sys.argv[1:4]
artigo = Document(caminho_artigo)
artigo.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
compositor = Composer(artigo)
compositor.append(Document(caminho_apendice))
compositor.save(caminho_saida)
print(f"composto {caminho_saida}")
